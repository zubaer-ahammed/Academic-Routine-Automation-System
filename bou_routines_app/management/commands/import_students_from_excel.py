from django.core.management.base import BaseCommand
from django.utils.dateparse import parse_date
from datetime import datetime
from bou_routines_app.models import Student, Centre, Semester
import openpyxl
import os


class Command(BaseCommand):
    help = 'Import students from Excel file (supports DRC, DUET, etc.)'

    def add_arguments(self, parser):
        parser.add_argument(
            'file_path',
            type=str,
            help='Path to the Excel file',
            default='students/Y1S1_course_enrollment_DRC.xlsx',
            nargs='?'
        )
        parser.add_argument(
            '--centre-code',
            type=str,
            help='Centre code (e.g., DRC, DUET). If not provided, will try to detect from filename.',
            default=None
        )
        parser.add_argument(
            '--centre-name',
            type=str,
            help='Centre name (e.g., "Dhaka Regional Center (DRC)", "DUET Study Center (DUET)"). If not provided, will use default based on code.',
            default=None
        )

    def handle(self, *args, **options):
        file_path = options['file_path']
        centre_code = options['centre_code']
        centre_name = options['centre_name']
        
        if not os.path.exists(file_path):
            self.stdout.write(self.style.ERROR(f'File not found: {file_path}'))
            return
        
        # Auto-detect centre code from filename if not provided
        if not centre_code:
            filename = os.path.basename(file_path)
            if 'DRC' in filename.upper():
                centre_code = 'DRC'
                if not centre_name:
                    centre_name = 'Dhaka Regional Center (DRC)'
            elif 'DUET' in filename.upper():
                centre_code = 'DUET'
                if not centre_name:
                    centre_name = 'DUET Study Center (DUET)'
            else:
                self.stdout.write(self.style.ERROR('Could not detect centre from filename. Please provide --centre-code'))
                return
        
        # Set default centre name if not provided
        if not centre_name:
            if centre_code == 'DRC':
                centre_name = 'Dhaka Regional Center (DRC)'
            elif centre_code == 'DUET':
                centre_name = 'DUET Study Center (DUET)'
            else:
                centre_name = f'{centre_code} Study Center'
        
        # Load workbook
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
        except Exception as e:
            self.stdout.write(self.style.ERROR(f'Error loading Excel file: {e}'))
            return
        
        # Get or create Centre
        centre, created = Centre.objects.get_or_create(
            code=centre_code,
            defaults={
                'name': centre_name,
                'is_active': True
            }
        )
        if created:
            self.stdout.write(self.style.SUCCESS(f'Created centre: {centre.name}'))
        else:
            # Update name if it has changed
            if centre.name != centre_name:
                centre.name = centre_name
                centre.save()
                self.stdout.write(f'Updated centre name to: {centre.name}')
            else:
                self.stdout.write(f'Using existing centre: {centre.name}')
        
        # Find Y1S1 semester (prefer new curriculum if multiple exist)
        semesters = Semester.objects.filter(name='Y1S1')
        if not semesters.exists():
            self.stdout.write(self.style.ERROR('Y1S1 semester not found. Please create it first.'))
            return
        elif semesters.count() > 1:
            # If multiple, prefer new curriculum
            semester = semesters.filter(curriculum__code='NEW').first() or semesters.first()
            self.stdout.write(self.style.WARNING(f'Multiple Y1S1 semesters found. Using: {semester.name} (ID: {semester.id})'))
        else:
            semester = semesters.first()
        
        # Headers are in row 2 (index 2, 1-based)
        # Data starts from row 3 (index 3, 1-based)
        headers = {}
        header_row = ws[2]  # Row 2 (1-based index)
        
        # Map column letters to field names
        for cell in header_row:
            if cell.value:
                col_letter = cell.column_letter
                header_value = str(cell.value).strip()
                headers[col_letter] = header_value
                # Debug: print header mapping
                self.stdout.write(f'Column {col_letter}: {header_value}')
        
        # Find column indices
        column_map = {}
        for col_letter, header in headers.items():
            header_lower = header.lower()
            if 'bou id' in header_lower or 'id' in header_lower:
                column_map['id'] = col_letter
            elif 'name' in header_lower and 'father' not in header_lower and 'mother' not in header_lower:
                column_map['name'] = col_letter
            elif 'gendar' in header_lower or 'gender' in header_lower:
                column_map['gender'] = col_letter
            elif 'mobile' in header_lower or 'phone' in header_lower:
                column_map['phone'] = col_letter
            elif 'email' in header_lower:
                column_map['email'] = col_letter
            elif "father's name" in header_lower or "father name" in header_lower:
                column_map['father_name'] = col_letter
            elif "mother's name" in header_lower or "mother name" in header_lower:
                column_map['mother_name'] = col_letter
            elif 'date of birth' in header_lower or 'dob' in header_lower:
                column_map['date_of_birth'] = col_letter
            elif 'academic year' in header_lower or 'session' in header_lower:
                column_map['session'] = col_letter
        
        self.stdout.write(f'Column mapping: {column_map}')
        
        # Process data rows (starting from row 3)
        created_count = 0
        updated_count = 0
        error_count = 0
        
        for row_num in range(3, ws.max_row + 1):
            row = ws[row_num]
            
            # Get student ID (required field)
            if 'id' not in column_map:
                self.stdout.write(self.style.ERROR('Student ID column not found'))
                break
            
            student_id = row[openpyxl.utils.column_index_from_string(column_map['id']) - 1].value
            if not student_id:
                continue  # Skip empty rows
            
            student_id = str(student_id).strip()
            
            # Get other fields
            def get_cell_value(col_key):
                if col_key not in column_map:
                    return None
                col_idx = openpyxl.utils.column_index_from_string(column_map[col_key]) - 1
                value = row[col_idx].value
                return str(value).strip() if value is not None else None
            
            name = get_cell_value('name')
            gender = get_cell_value('gender')
            phone = get_cell_value('phone')
            email = get_cell_value('email')
            father_name = get_cell_value('father_name')
            mother_name = get_cell_value('mother_name')
            dob_str = get_cell_value('date_of_birth')
            session = get_cell_value('session')
            
            # Parse date of birth (format: d-m-y, e.g., 28-04-2006)
            date_of_birth = None
            if dob_str:
                try:
                    # Parse d-m-y format
                    date_of_birth = datetime.strptime(dob_str, '%d-%m-%Y').date()
                except ValueError:
                    try:
                        # Try alternative formats
                        date_of_birth = datetime.strptime(dob_str, '%d/%m/%Y').date()
                    except ValueError:
                        self.stdout.write(self.style.WARNING(f'Could not parse date: {dob_str} for student {student_id}'))
            
            # Normalize gender
            if gender:
                gender = gender.strip().title()
                if gender not in ['Male', 'Female', 'Other']:
                    # Try to normalize
                    if 'male' in gender.lower():
                        gender = 'Male'
                    elif 'female' in gender.lower():
                        gender = 'Female'
                    else:
                        gender = 'Other'
            
            # Create or update student
            try:
                student, created = Student.objects.update_or_create(
                    id=student_id,
                    defaults={
                        'name': name or '',
                        'gender': gender,
                        'phone': phone,
                        'email': email,
                        'father_name': father_name,
                        'mother_name': mother_name,
                        'date_of_birth': date_of_birth,
                        'session': session or '',
                        'centre': centre,
                    }
                )
                
                # Add semester if not already added
                if semester not in student.semesters.all():
                    student.semesters.add(semester)
                
                if created:
                    created_count += 1
                    self.stdout.write(f'Created: {student_id} - {name}')
                else:
                    updated_count += 1
                    self.stdout.write(f'Updated: {student_id} - {name}')
                    
            except Exception as e:
                error_count += 1
                self.stdout.write(self.style.ERROR(f'Error processing {student_id}: {e}'))
        
        # Summary
        self.stdout.write(self.style.SUCCESS(
            f'\nImport complete!\n'
            f'Created: {created_count}\n'
            f'Updated: {updated_count}\n'
            f'Errors: {error_count}'
        ))

